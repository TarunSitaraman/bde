"""Build the assignment document (.docx) and presentation (.pptx) from results/*.json."""
import json
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from pptx import Presentation
from pptx.util import Inches as PI, Pt as PP
from pptx.dml.color import RGBColor as PC

URL = "https://github.com/TarunSitaraman/bde/tree/claude/ecstatic-johnson-b4ik3x"
C = json.load(open("results/classification_results.json")); K = json.load(open("results/clustering_results.json"))
M = C["_meta"]; best = M["best_model"]; B = C[best]; cm = B["confusion"]
base = C["Baseline (last week's class)"]["accuracy"]
prof = {r["demand_tier"]: r for r in K["profile"]}
sel = K["k_selection"]
pct = lambda x: f"{x*100:.1f}%"
inr = lambda x: f"₹{x/1e3:,.0f}k"
code_cls = open("src/classification.py").read(); code_clu = open("src/clustering.py").read()
def snippet(src, start, end):
    a = src.index(start); return src[a:src.index(end, a)].rstrip()
CLS_SNIP = snippet(code_cls, "w = Window", "# ---------- 3.")
CLS_SNIP2 = snippet(code_cls, "models = {", "results, fitted")
CLU_SNIP = snippet(code_clu, "# ---------- ATM-level", "# ---------- Choose k")
CLU_SNIP2 = snippet(code_clu, "km = KMeans(k=3", "out = km.transform")

# ============================== DOCX ==============================
d = Document()
st = d.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
def H(t, l=1): d.add_heading(t, l)
def P(t, b=False):
    p = d.add_paragraph(); r = p.add_run(t); r.bold = b; return p
def BL(items):
    for t in items: d.add_paragraph(t, style="List Bullet")
def CODE(t):
    p = d.add_paragraph(); r = p.add_run(t); r.font.name = "Consolas"; r.font.size = Pt(8)
def TABLE(rows, header=True):
    t = d.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Light Grid Accent 1"
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = t.cell(i, j); c.text = str(v)
            for p in c.paragraphs:
                for r in p.runs: r.font.size = Pt(9.5); r.bold = (i == 0 and header)
    d.add_paragraph()

d.add_heading("Predictive ATM Cash Management using Big Data Analytics", 0)
P("BDE Assignment – Team Document (Classification + Clustering)", True)
P("Technology: Apache Spark 4 (PySpark MLlib), Python, local mode • Data: synthetic ATM transaction data (219,300 daily records, 300 ATMs, 2 years)")
P(f"Implementation URL: {URL}")

H("0. Data and Big-Data Pipeline (common to both problems)")
P("No public per-ATM transaction dataset was available, so a synthetic dataset was generated (src/generate_data.py, fixed seed 42). "
  "It mimics real ATM behaviour: seven location types with different base demand, weekend effects, payday spikes (days 28–3), festival "
  "surges, 10% growth over two years, log-normal noise, and random ATM outages. Every conclusion below therefore describes the pipeline and "
  "method on this data; absolute accuracy on a real bank's data should be re-measured.")
TABLE([["Attribute", "Description"],
       ["atm_id, location_type", "ATM identifier; Branch / Mall / Transit_Hub / Residential / Highway / Hospital / Rural"],
       ["date", "Day of transaction (2024-01-01 … 2025-12-31)"],
       ["total_withdrawal, tx_count, failed_tx", "Rupees dispensed, successful transactions, failed transactions that day"],
       ["is_weekend, is_payday, is_festival", "Calendar flags"]])
P("Spark reads the daily CSV, aggregates with group-by / window functions (weekly roll-ups, lags, rolling statistics for classification; "
  "ATM-level behavioural profiles for clustering) and trains MLlib models — the same code scales from this laptop-sized sample to a cluster "
  "handling billions of rows.")

# ---------------- Classification ----------------
H("1. Classification Problem – High vs Low Cash Demand")
H("1.1 Problem Statement", 2)
P("Predict, at the start of each week, whether an ATM will have HIGH or LOW cash demand in that week, using only its past transaction and "
  "withdrawal patterns plus calendar information known in advance. High demand = weekly withdrawals above "
  f"₹{M['threshold_inr_per_week']:,.0f} (the median of the training weeks). Business value: ATMs flagged HIGH get larger or earlier cash "
  "replenishment, preventing stock-outs, while LOW ATMs avoid idle cash that incurs holding cost.")
H("1.2 Model and Justification", 2)
P(f"Three supervised binary classifiers were trained and compared on a time-based split (train: weeks before {M['split']}, "
  f"{M['train_rows']:,} ATM-weeks; test: later weeks, {M['test_rows']:,} ATM-weeks), which mimics real deployment by never training on the future:")
BL(["Logistic Regression – fast, interpretable, a strong baseline for demand that is roughly log-linear in past demand.",
    "Random Forest (150 trees) – captures non-linear interactions and gives feature importance.",
    "Gradient-Boosted Trees (60 iterations) – typically the most accurate tabular learner."])
P(f"Selected model: {best} (highest AUC on held-out future weeks). The three models are within about one percentage point, which is "
  "expected because demand persistence dominates; the simplest model that matches the complex ones is preferred (cheap to retrain "
  "for thousands of ATMs, explainable to operations teams). Features (13 numeric + location type): lagged weekly withdrawals (lag1, lag2), "
  "4-week rolling mean and std, last week’s transaction count, average ticket size, failure rate, weekend share, outage days, "
  "payday days and festival days in the target week, and seasonal sine/cosine of week-of-year. All lag features use only past weeks (no leakage).")
H("1.3 Coding", 2)
P("Full code: src/classification.py. Key parts (feature engineering with Spark window functions, and model definitions):")
CODE(CLS_SNIP); CODE(CLS_SNIP2)
H("1.4 Results", 2)
rows = [["Model", "Accuracy", "F1", "AUC"]]
for n in ["Logistic Regression", "Random Forest", "Gradient-Boosted Trees"]:
    rows.append([n, pct(C[n]["accuracy"]), f"{C[n]['f1']:.3f}", f"{C[n]['auc']:.3f}"])
rows.append(["Baseline: repeat last week’s class", pct(base), f"{C['Baseline (last week' + chr(39) + 's class)']['f1']:.3f}", "–"])
TABLE(rows)
P(f"Confusion matrix of {best} on {M['test_rows']:,} unseen ATM-weeks: TN={cm['TN']:,}, FP={cm['FP']:,}, FN={cm['FN']:,}, TP={cm['TP']:,}. "
  f"For the HIGH class, precision = {pct(M['high_class_precision'])} and recall = {pct(M['high_class_recall'])}.")
d.add_picture("results/classification_model_comparison.png", width=Inches(4.6))
d.add_picture("results/classification_plots.png", width=Inches(6.4))
H("1.5 Inference", 2)
BL([f"{best} reaches {pct(B['accuracy'])} accuracy and AUC {B['auc']:.3f} on future weeks, versus {pct(base)} for the naive ‘same as last week’ rule — "
    f"a {(B['accuracy']-base)*100:.1f}-point gain, largely from catching ATMs whose demand crosses the threshold (e.g. payday and festival weeks).",
    "Recent withdrawals dominate: the 4-week rolling mean, last week’s withdrawals/transactions and the week before account for ~88% of "
    "Random Forest importance; payday days come next, then ATM location type. Calendar knowledge adds measurable but smaller lift.",
    f"Operational risk: {cm['FN']} high-demand ATM-weeks ({pct(1-M['high_class_recall'])}) were predicted LOW — the costly stock-out error. "
    "Lowering the decision threshold (e.g. to 0.4) would trade some false alarms for fewer missed high-demand weeks.",
    "Caveats: results are on synthetic data with clean structure; on real data expect lower accuracy and re-tune. Trees did not beat logistic "
    "regression here because the generated demand is multiplicative; real data with outages/events may favour boosted trees."])
H("1.6 URL of Implementation", 2)
P(URL + "  (src/classification.py; outputs in results/)")

# ---------------- Clustering ----------------
H("2. Clustering Problem – ATM Demand Tiers")
H("2.1 Problem Statement", 2)
P("Group the 300 ATMs into LOW-, MEDIUM- and HIGH-demand categories purely from their withdrawal behaviour (no labels), so that "
  "replenishment frequency, cash-in-transit routes and cash limits can be set per tier instead of per machine.")
H("2.2 Model and Justification", 2)
P("Model: K-Means with k = 3 on standardised, log-transformed volume features (average daily withdrawal, 90th-percentile daily withdrawal, "
  "average daily transactions). Justification:")
BL(["Unsupervised, distance-based and scalable (Spark MLlib K-Means runs in parallel over millions of ATM profiles).",
    "k = 3 is dictated by the business tiers (Low/Medium/High). It is supported statistically: the elbow flattens after k=3–4 and the "
    f"silhouette at k=3 is {sel['3']['silhouette']:.2f} (k=2 scores {sel['2']['silhouette']:.2f}, but two groups cannot express a medium tier).",
    "Log transform is applied because demand is right-skewed; without it, K-Means isolated just 5 very busy ATMs as ‘High’ and put 193 in ‘Low’ — useless for planning.",
    f"Compared with a Gaussian Mixture (silhouette {K['algorithm_comparison_silhouette']['Gaussian Mixture (k=3)']:.2f}), K-Means "
    f"({K['algorithm_comparison_silhouette']['K-Means (k=3)']:.2f}) gives better-separated, simpler clusters. Other behavioural metrics "
    "(variability, weekend ratio, payday uplift, failure rate) are used only to profile the tiers, not to form them."])
H("2.3 Coding", 2)
P("Full code: src/clustering.py. Key parts:")
CODE(CLU_SNIP); CODE(CLU_SNIP2)
H("2.4 Results", 2)
rows = [["Tier", "ATMs", "Avg daily withdrawal", "P90 daily withdrawal", "Avg daily txns", "Payday uplift"]]
for t in ["Low", "Medium", "High"]:
    r = prof[t]; rows.append([t, r["n_atms"], inr(r["avg_daily_wd"]), inr(r["p90_daily_wd"]), f"{r['avg_daily_tx']:.0f}", f"{r['payday_uplift']:.2f}x"])
TABLE(rows)
P(f"Silhouette (k=3): {sel['3']['silhouette']:.3f}. Location-type composition of the tiers:")
mix = K["location_mix"]
TABLE([["Location", "Low", "Medium", "High"]] + [[m["location_type"], m["Low"], m["Medium"], m["High"]] for m in mix])
d.add_picture("results/clustering_plots.png", width=Inches(6.5))
d.add_picture("results/clustering_boxplot.png", width=Inches(4.2))
H("2.5 Inference", 2)
BL([f"The High tier withdraws about {prof['High']['avg_daily_wd']/prof['Low']['avg_daily_wd']:.1f}× the cash of the Low tier "
    f"({inr(prof['High']['avg_daily_wd'])} vs {inr(prof['Low']['avg_daily_wd'])} per day), so a single replenishment schedule would either starve the High tier or over-stock the Low tier.",
    "Tiers align with location: Malls, Transit hubs and Branches are mostly High; Residential areas are mostly Medium; Rural, Highway and Hospital ATMs are mostly Low. New ATM sites can be pre-assigned a tier from their location type.",
    "Low-demand ATMs have the strongest payday uplift (1.35×) — small machines run dry around month-end even though they are quiet otherwise; they need payday top-ups despite the Low label.",
    "Suggested policy: High – daily monitoring and frequent loading; Medium – every 2–3 days; Low – weekly plus payday top-up. The classifier in Part 1 then adjusts each ATM week by week.",
    "Caveat: demand is a continuum, so tier boundaries are partly a business convention (k=2 scores higher silhouette); ATMs near a boundary can move tiers between periods."])
H("2.6 URL of Implementation", 2)
P(URL + "  (src/clustering.py; outputs in results/atm_clusters.csv)")
H("3. How to reproduce", 1)
CODE("pip install -r requirements.txt\npython src/generate_data.py\npython src/classification.py\npython src/clustering.py\npython src/build_docs.py")
d.save("docs/BDE_Assignment_Report.docx")

# ============================== PPTX ==============================
pr = Presentation(); pr.slide_width, pr.slide_height = PI(13.333), PI(7.5)
NAVY, BLUE = PC(0x0B, 0x25, 0x4A), PC(0x1F, 0x6F, 0xEB)
def slide(title, bullets=None, img=None, img_w=6.2, code=None, sub=None):
    s = pr.slides.add_slide(pr.slide_layouts[6])
    bar = s.shapes.add_shape(1, 0, 0, pr.slide_width, PI(1.0)); bar.fill.solid(); bar.fill.fore_color.rgb = NAVY; bar.line.fill.background()
    tb = s.shapes.add_textbox(PI(.5), PI(.18), PI(12.3), PI(.7)); r = tb.text_frame.paragraphs[0].add_run(); r.text = title
    r.font.size = PP(28); r.font.bold = True; r.font.color.rgb = PC(255, 255, 255)
    width = 6.4 if (img or code) else 12.3
    if bullets:
        tf = s.shapes.add_textbox(PI(.5), PI(1.3), PI(width), PI(5.6)).text_frame; tf.word_wrap = True
        for i, b in enumerate(bullets):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.text = "• " + b; p.font.size = PP(17); p.space_after = PP(9)
    if img: s.shapes.add_picture(img, PI(7.0 if bullets else 1.5), PI(1.4), width=PI(img_w if bullets else 10))
    if code:
        tb = s.shapes.add_textbox(PI(7.0), PI(1.3), PI(6), PI(5.8)); tb.text_frame.word_wrap = True
        tb.fill.solid(); tb.fill.fore_color.rgb = PC(0xF3, 0xF4, 0xF6)
        for i, ln in enumerate(code.split("\n")):
            p = tb.text_frame.paragraphs[0] if i == 0 else tb.text_frame.add_paragraph(); p.text = ln; p.font.size = PP(12); p.font.name = "Consolas"
    return s

s = slide("Predictive ATM Cash Management", [ "using Big Data Analytics (Apache Spark / PySpark MLlib)", "Classification: High vs Low demand  •  Clustering: Low / Medium / High ATM tiers", f"Code: {URL}"])
slide("Problem & Data", ["Cash-outs hurt customers; over-stocking ties up capital → forecast demand per ATM",
      "Data: 219,300 daily records • 300 ATMs • 2 years • 7 location types (synthetic, seed 42)",
      "Fields: withdrawal ₹, txn count, failures, weekend / payday / festival flags",
      "Spark: group-by roll-ups, window lags, rolling stats, MLlib pipelines (scales to billions of rows)",
      "Task 1 – classify next-week demand High/Low;  Task 2 – cluster ATMs into 3 tiers"])
slide("Model Selection & Justification (1 mark)", [
      "Classification: compared Logistic Regression, Random Forest, GBT on time-based split (no future leakage)",
      f"Winner: {best} – best AUC ({B['auc']:.3f}); simple, fast, explainable; trees ≈ equal",
      "Clustering: K-Means k=3 on log-scaled volume features",
      f"k=3 = business tiers; silhouette {sel['3']['silhouette']:.2f}; beats Gaussian Mixture ({K['algorithm_comparison_silhouette']['Gaussian Mixture (k=3)']:.2f})",
      "Log transform fixes skew: raw K-Means gave tiers of 193 / 102 / 5 ATMs"])
slide("Coding & Implementation – Classification (2 marks)", [
      "Weekly aggregation of daily data in Spark", "Lag-1/2, 4-wk rolling mean/std via Window functions (past only)",
      "Calendar features: payday days, festival days, week-of-year sin/cos", f"Label threshold = train median (₹{M['threshold_inr_per_week']/1e6:.2f}M/week)",
      "Spark ML Pipeline: StringIndexer → OneHot → Assembler → Scaler → classifier"], code="w = Window.partitionBy('atm_id').orderBy('week')\nw4 = w.rowsBetween(-4, -1)\n\nfeat = (weekly\n  .withColumn('lag1_wd', F.lag('wd',1).over(w))\n  .withColumn('roll4_wd', F.avg('wd').over(w4))\n  .withColumn('woy_sin', F.sin(2*pi*woy/52)))\n\nthr = train.approxQuantile('wd',[.5],.001)[0]\nlabel = (F.col('wd') > thr).cast('double')\n\nPipeline(stages=[indexer, onehot,\n   assembler, scaler, LogisticRegression()])")
slide("Coding & Implementation – Clustering", [
      "ATM-level profile via Spark groupBy (one row per ATM)", "Features: avg, P90 daily withdrawal, avg daily txns → log1p → StandardScaler",
      "Elbow + silhouette for k=2..7; final KMeans k=3", "Tiers named by ranking cluster mean withdrawal",
      "Other metrics (CV, payday uplift…) used only to profile"], code="prof = daily.groupBy('atm_id').agg(\n  avg('total_withdrawal'),\n  percentile_approx(..., 0.9),\n  avg('tx_count'))\n\n# log1p -> StandardScaler\nfor k in range(2, 8):\n    KMeans(k=k).fit(data)  # elbow+silhouette\n\nkm = KMeans(k=3, seed=11).fit(data)\n# rank clusters by mean withdrawal\n# -> Low / Medium / High")
rows = [f"{n}: acc {pct(C[n]['accuracy'])}, AUC {C[n]['auc']:.3f}" for n in ["Logistic Regression", "Random Forest", "Gradient-Boosted Trees"]]
slide("Results – Classification (1 mark)", rows + [f"Baseline (repeat last week): {pct(base)}", f"High-class recall {pct(M['high_class_recall'])}, precision {pct(M['high_class_precision'])}"], img="results/classification_plots.png", img_w=6.1)
slide("Results – Clustering", [f"Low: {prof['Low']['n_atms']} ATMs, {inr(prof['Low']['avg_daily_wd'])}/day", f"Medium: {prof['Medium']['n_atms']} ATMs, {inr(prof['Medium']['avg_daily_wd'])}/day",
      f"High: {prof['High']['n_atms']} ATMs, {inr(prof['High']['avg_daily_wd'])}/day", f"Silhouette {sel['3']['silhouette']:.2f}"], img="results/clustering_plots.png", img_w=6.1)
slide("Analysis & Inference", [
      f"Classifier beats naive baseline by {(B['accuracy']-base)*100:.1f} pts; recent withdrawals drive ~88% of importance",
      f"Missed high-demand weeks = {pct(1-M['high_class_recall'])} → tune threshold to cut stock-out risk",
      f"High tier needs ~{prof['High']['avg_daily_wd']/prof['Low']['avg_daily_wd']:.1f}× the cash of Low; malls/transit/branches are mostly High",
      "Low-demand ATMs show strongest payday spike → add month-end top-ups",
      "Combined policy: tier sets baseline schedule, weekly classifier adjusts it",
      "Limitation: synthetic data – re-validate on real bank data"])
slide("Thank You", [f"Code, results and report: {URL}"])
pr.save("docs/BDE_Assignment_Presentation.pptx")
print("built docs")
